from typing import Optional, Union, Tuple
from pathlib import Path
from io import BytesIO
from docx.document import Document
from docx.text.paragraph import Paragraph
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

from app.schemas.ast import ASTNode


def render_figure(
    doc: Document,
    image_path_or_stream: Optional[Union[str, Path, BytesIO]] = None,
    caption: Optional[str] = None,
    figure_num: Optional[str] = None,
    width_cm: Optional[float] = None,
    node: Optional[ASTNode] = None
) -> Tuple[Paragraph, Optional[Paragraph]]:
    """
    Рендерит рисунок / иллюстрацию по ГОСТ 7.32-2017:
    - Выравнивание рисунка: По центру (CENTER)
    - Отступы рисунка: Space before 6pt, Space after 6pt, keep_with_next = True
    - Подпись ПОД рисунком: Рисунок X.Y — Название рисунка (без точки в конце)
    - Выравнивание подписи: По центру (CENTER)
    - Шрифт подписи: Times New Roman 12pt, Space after 12pt
    """
    fig_caption = caption
    if node is not None and fig_caption is None:
        fig_caption = node.caption

    # 1. Параграф с изображением
    p_img = doc.add_paragraph()
    p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_img.paragraph_format.first_line_indent = Cm(0)
    p_img.paragraph_format.space_before = Pt(6)
    p_img.paragraph_format.space_after = Pt(6)
    p_img.paragraph_format.keep_with_next = True

    if image_path_or_stream is not None:
        run_img = p_img.add_run()
        if width_cm:
            run_img.add_picture(image_path_or_stream, width=Cm(width_cm))
        else:
            run_img.add_picture(image_path_or_stream)

    # 2. Подпись под рисунком
    p_cap: Optional[Paragraph] = None
    if fig_caption:
        clean_cap = fig_caption.strip()
        if clean_cap.endswith("."):
            clean_cap = clean_cap[:-1].strip()

        if not clean_cap.startswith("Рисунок"):
            prefix = f"Рисунок {figure_num} — " if figure_num else "Рисунок — "
            formatted_caption = f"{prefix}{clean_cap}"
        else:
            formatted_caption = clean_cap

        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.paragraph_format.first_line_indent = Cm(0)
        p_cap.paragraph_format.space_before = Pt(6)
        p_cap.paragraph_format.space_after = Pt(12)

        run_cap = p_cap.add_run(formatted_caption)
        run_cap.font.name = "Times New Roman"
        run_cap.font.size = Pt(12)
        run_cap.font.color.rgb = RGBColor(0, 0, 0)

    return p_img, p_cap
