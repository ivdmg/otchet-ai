from typing import Optional
import docx
from docx.document import Document
from docx.table import Table
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from app.schemas.ast import ASTNode


def remove_table_borders(table: Table) -> None:
    """
    Убирает все границы у таблицы и ее ячеек (val='none').
    """
    tbl_pr = table._tbl.tblPr
    tbl_borders = OxmlElement('w:tblBorders')
    for border_name in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
        border = OxmlElement(f'w:{border_name}')
        border.set(qn('w:val'), 'none')
        tbl_borders.append(border)
    tbl_pr.append(tbl_borders)

    for row in table.rows:
        for cell in row.cells:
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_borders = OxmlElement('w:tcBorders')
            for border_name in ['top', 'left', 'bottom', 'right']:
                border = OxmlElement(f'w:{border_name}')
                border.set(qn('w:val'), 'none')
                tc_borders.append(border)
            tc_pr.append(tc_borders)


def render_formula(
    doc: Document,
    formula_text: Optional[str] = None,
    formula_num: Optional[str] = None,
    node: Optional[ASTNode] = None
) -> Table:
    """
    Рендерит математическую формулу по ГОСТ 7.32-2017:
    - 1-строчная бескаркасная таблица 100% ширины
    - Ячейка 0: Формула по центру (CENTER), Times New Roman 14pt Italic
    - Ячейка 1: Номер формулы справа (RIGHT) в круглых скобках (X.Y)
    - Отступы До/После: 6 pt
    """
    f_text = formula_text
    f_num = formula_num

    if node is not None:
        if f_text is None and node.formula_text:
            f_text = node.formula_text
        if f_num is None and node.formula_number:
            f_num = node.formula_number

    f_text = f_text or ""
    f_num = f_num or ""

    table = doc.add_table(rows=1, cols=2)
    table.autofit = False

    table.columns[0].width = Cm(14.0)
    table.columns[1].width = Cm(3.0)

    remove_table_borders(table)

    # Ячейка 0 (Текст формулы)
    p0 = table.cell(0, 0).paragraphs[0]
    p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p0.paragraph_format.first_line_indent = Cm(0)
    p0.paragraph_format.space_before = Pt(6)
    p0.paragraph_format.space_after = Pt(6)

    run0 = p0.add_run(f_text)
    run0.font.name = "Times New Roman"
    run0.font.size = Pt(14)
    run0.font.italic = True
    run0.font.color.rgb = RGBColor(0, 0, 0)

    # Ячейка 1 (Номер формулы)
    p1 = table.cell(0, 1).paragraphs[0]
    p1.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p1.paragraph_format.first_line_indent = Cm(0)
    p1.paragraph_format.space_before = Pt(6)
    p1.paragraph_format.space_after = Pt(6)

    if f_num:
        formatted_num = f_num.strip()
        if not (formatted_num.startswith("(") and formatted_num.endswith(")")):
            formatted_num = f"({formatted_num})"

        run1 = p1.add_run(formatted_num)
        run1.font.name = "Times New Roman"
        run1.font.size = Pt(14)
        run1.font.italic = False
        run1.font.color.rgb = RGBColor(0, 0, 0)

    return table
