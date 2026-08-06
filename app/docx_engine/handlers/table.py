from typing import Optional, List, Tuple
import docx
from docx.document import Document
from docx.table import Table
from docx.text.paragraph import Paragraph
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from app.schemas.ast import ASTNode


def set_table_borders(table: Table) -> None:
    """
    Устанавливает сплошные черные границы 0.5pt в OpenXML таблицы.
    """
    tbl_pr = table._tbl.tblPr
    tbl_borders = OxmlElement('w:tblBorders')
    for border_name in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
        border = OxmlElement(f'w:{border_name}')
        border.set(qn('w:val'), 'single')
        border.set(qn('w:sz'), '4')  # 0.5 pt = 4/8 pt
        border.set(qn('w:space'), '0')
        border.set(qn('w:color'), '000000')
        tbl_borders.append(border)
    tbl_pr.append(tbl_borders)


def render_table(
    doc: Document,
    headers: Optional[List[str]] = None,
    rows: Optional[List[List[str]]] = None,
    caption: Optional[str] = None,
    table_num: Optional[str] = None,
    node: Optional[ASTNode] = None
) -> Tuple[Optional[Paragraph], Table]:
    """
    Рендерит таблицу по ГОСТ 7.32-2017:
    - Подпись НАД таблицей (Таблица X.Y — Название таблицы)
    - Повтор шапки на каждой странице (w:tblHeader)
    - Запрет разрыва строк (w:cantSplit)
    - Границы 0.5pt (w:tblBorders)
    - Шрифт внутри: Times New Roman 12pt, одинарный интервал, space 0pt
    - Шапка: по центру, bold
    - Данные: по левому краю
    """
    tbl_headers = headers
    tbl_rows = rows
    tbl_caption = caption

    if node is not None:
        if tbl_headers is None and node.headers:
            tbl_headers = node.headers
        if tbl_rows is None and node.rows:
            tbl_rows = node.rows
        if tbl_caption is None and node.caption:
            tbl_caption = node.caption

    tbl_headers = tbl_headers or []
    tbl_rows = tbl_rows or []

    # 1. Подпись таблицы
    caption_p: Optional[Paragraph] = None
    if tbl_caption:
        clean_cap = tbl_caption.strip()
        if clean_cap.endswith("."):
            clean_cap = clean_cap[:-1].strip()

        if not clean_cap.startswith("Таблица"):
            prefix = f"Таблица {table_num} — " if table_num else "Таблица — "
            formatted_caption = f"{prefix}{clean_cap}"
        else:
            formatted_caption = clean_cap

        caption_p = doc.add_paragraph()
        caption_p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        caption_p.paragraph_format.first_line_indent = Cm(1.25)
        caption_p.paragraph_format.space_before = Pt(6)
        caption_p.paragraph_format.space_after = Pt(6)
        caption_p.paragraph_format.keep_with_next = True

        run = caption_p.add_run(formatted_caption)
        run.font.name = "Times New Roman"
        run.font.size = Pt(12)
        run.font.color.rgb = RGBColor(0, 0, 0)

    # 2. Размеры таблицы
    num_cols = max(len(tbl_headers), max((len(r) for r in tbl_rows), default=0))
    if num_cols == 0:
        num_cols = 1

    num_rows = (1 if tbl_headers else 0) + len(tbl_rows)
    table = doc.add_table(rows=num_rows, cols=num_cols)
    set_table_borders(table)

    current_row_idx = 0

    # 3. Оформление шапки (Header Row)
    if tbl_headers:
        hdr_row = table.rows[0]
        tr_pr = hdr_row._tr.get_or_add_trPr()
        tr_pr.append(OxmlElement('w:tblHeader'))

        for i, header_text in enumerate(tbl_headers):
            cell = hdr_row.cells[i]
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.line_spacing = 1.0
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.first_line_indent = Cm(0)

            run = p.add_run(header_text)
            run.font.name = "Times New Roman"
            run.font.size = Pt(12)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0, 0, 0)

        current_row_idx = 1

    # 4. Оформление ячеек данных
    for row_data in tbl_rows:
        row = table.rows[current_row_idx]
        current_row_idx += 1

        for i, val in enumerate(row_data):
            if i < num_cols:
                cell = row.cells[i]
                p = cell.paragraphs[0]
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                p.paragraph_format.line_spacing = 1.0
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.first_line_indent = Cm(0)

                run = p.add_run(str(val))
                run.font.name = "Times New Roman"
                run.font.size = Pt(12)
                run.font.bold = False
                run.font.color.rgb = RGBColor(0, 0, 0)

    # 5. Применение w:cantSplit для всех строк таблицы
    for row in table.rows:
        tr_pr = row._tr.get_or_add_trPr()
        tr_pr.append(OxmlElement('w:cantSplit'))

    return caption_p, table
