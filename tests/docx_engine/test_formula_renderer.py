import pytest
import docx
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

from app.docx_engine.handlers.formula import render_formula
from app.schemas.ast import ASTNode, NodeType


def test_render_formula_structure():
    """Проверка создания 2-колонковой бескаркасной таблицы с формулой и номером."""
    doc = docx.Document()
    formula_text = "S = \\pi r^2"
    formula_num = "1.1"

    table = render_formula(
        doc,
        formula_text=formula_text,
        formula_num=formula_num
    )

    assert table is not None
    assert len(table.rows) == 1
    assert len(table.columns) == 2

    # Ячейка 0 (Формула по центру, курсив)
    p0 = table.cell(0, 0).paragraphs[0]
    assert p0.alignment == WD_ALIGN_PARAGRAPH.CENTER
    assert p0.text == formula_text
    assert len(p0.runs) > 0
    assert p0.runs[0].font.name == "Times New Roman"
    assert p0.runs[0].font.size == Pt(14)
    assert p0.runs[0].font.italic is True

    # Ячейка 1 (Номер справа в круглых скобках)
    p1 = table.cell(0, 1).paragraphs[0]
    assert p1.alignment == WD_ALIGN_PARAGRAPH.RIGHT
    assert p1.text == "(1.1)"
    assert len(p1.runs) > 0
    assert p1.runs[0].font.name == "Times New Roman"
    assert p1.runs[0].font.size == Pt(14)


def test_render_formula_borderless_xml():
    """Проверка отсутствия границ в OpenXML таблицы формулы (val='none')."""
    doc = docx.Document()
    table = render_formula(doc, formula_text="A = B + C", formula_num="2.1")

    tbl_xml = table._tbl.xml
    # Таблица или ячейки должны содержать val="none"
    assert 'w:val="none"' in tbl_xml or 'val="none"' in tbl_xml


def test_render_formula_from_ast_node():
    """Проверка создания формулы из узла ASTNode."""
    doc = docx.Document()
    node = ASTNode(
        id="f_node_1",
        type=NodeType.FORMULA,
        formula_text="F = m \\cdot a",
        formula_number="(3.2)"
    )

    table = render_formula(doc, node=node)
    assert table.cell(0, 0).paragraphs[0].text == "F = m \\cdot a"
    assert table.cell(0, 1).paragraphs[0].text == "(3.2)"
