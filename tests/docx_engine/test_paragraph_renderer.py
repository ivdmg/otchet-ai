import pytest
import docx
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

from app.docx_engine.compiler import DocxCompiler
from app.docx_engine.handlers.paragraph import render_paragraph
from app.schemas.ast import ASTNode, NodeType


def test_render_paragraph_formatting():
    """Проверка правильности стилизации абзаца по ГОСТ 7.32-2017."""
    doc = docx.Document()
    sample_text = "Тестовый параграф основного текста для проверки стандартов ГОСТ."
    p = render_paragraph(doc, text=sample_text)

    assert p is not None
    assert p.text == sample_text

    # Выравнивание по ширине
    assert p.alignment == WD_ALIGN_PARAGRAPH.JUSTIFY

    # Межстрочный полуторный (1.5)
    assert p.paragraph_format.line_spacing == 1.5

    # Отступы До и После = 0pt
    assert p.paragraph_format.space_before == Pt(0)
    assert p.paragraph_format.space_after == Pt(0)

    # Красная строка 1.25 см
    assert p.paragraph_format.first_line_indent is not None
    assert abs(p.paragraph_format.first_line_indent.cm - 1.25) < 0.05

    # Проверка свойств шрифта
    assert len(p.runs) == 1
    run = p.runs[0]
    assert run.font.name == "Times New Roman"
    assert run.font.size == Pt(14)
    assert run.font.color.rgb == RGBColor(0, 0, 0)
    assert run.font.bold is False or run.font.bold is None


def test_render_paragraph_openxml():
    """Проверка наличия OpenXML атрибута w:firstLine в свойствах абзаца."""
    doc = docx.Document()
    p = render_paragraph(doc, text="Проверка XML")
    
    xml_str = p._p.xml
    assert "w:ind" in xml_str
    assert "w:firstLine" in xml_str


def test_render_paragraph_from_ast_node():
    """Проверка рендеринга параграфа из ASTNode."""
    doc = docx.Document()
    node = ASTNode(
        id="node_p1",
        type=NodeType.PARAGRAPH,
        text="Параграф из AST дерева."
    )
    p = render_paragraph(doc, node=node)
    assert p.text == "Параграф из AST дерева."
    assert p.alignment == WD_ALIGN_PARAGRAPH.JUSTIFY
