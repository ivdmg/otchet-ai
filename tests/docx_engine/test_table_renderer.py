import pytest
import docx
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

from app.docx_engine.handlers.table import render_table
from app.schemas.ast import ASTNode, NodeType


def test_render_table_structure_and_caption():
    """Проверка структуры таблицы, подписи НАД таблицей и формата шрифта."""
    doc = docx.Document()
    headers = ["Показатель", "2024 г.", "2025 г."]
    rows = [
        ["Объем ипотеки, млрд руб.", "1200", "1450"],
        ["Средняя ставка, %", "8.5", "9.2"]
    ]
    caption_text = "Динамика показателей ипотечного рынка."

    caption_p, table = render_table(
        doc,
        headers=headers,
        rows=rows,
        caption=caption_text,
        table_num="1.1"
    )

    # Проверка подписи
    assert caption_p is not None
    # Точка в конце подписи должна быть удалена
    assert caption_p.text == "Таблица 1.1 — Динамика показателей ипотечного рынка"
    assert caption_p.paragraph_format.keep_with_next is True

    # Проверка размера таблицы
    assert len(table.rows) == 3  # 1 header + 2 data rows
    assert len(table.columns) == 3

    # Проверка ячеек шапки (по центру, bold)
    hdr_cells = table.rows[0].cells
    for i, title in enumerate(headers):
        assert hdr_cells[i].text == title
        p = hdr_cells[i].paragraphs[0]
        assert p.alignment == WD_ALIGN_PARAGRAPH.CENTER
        assert len(p.runs) > 0
        assert p.runs[0].font.bold is True
        assert p.runs[0].font.name == "Times New Roman"
        assert p.runs[0].font.size == Pt(12)

    # Проверка интервала внутри ячеек (одинарный, space 0)
    data_cell = table.rows[1].cells[0]
    p_data = data_cell.paragraphs[0]
    assert p_data.paragraph_format.line_spacing == 1.0
    assert p_data.paragraph_format.space_before == Pt(0)
    assert p_data.paragraph_format.space_after == Pt(0)


def test_render_table_openxml_properties():
    """Проверка наличия w:tblHeader, w:cantSplit и w:tblBorders в OpenXML таблицы."""
    doc = docx.Document()
    headers = ["Колонка 1", "Колонка 2"]
    rows = [["Данные 1", "Данные 2"]]

    _, table = render_table(doc, headers=headers, rows=rows)

    # Проверка w:tblHeader на 1-й строке
    hdr_tr_xml = table.rows[0]._tr.xml
    assert "w:tblHeader" in hdr_tr_xml

    # Проверка w:cantSplit на всех строках
    for row in table.rows:
        tr_xml = row._tr.xml
        assert "w:cantSplit" in tr_xml

    # Проверка границ таблицы w:tblBorders
    tbl_xml = table._tbl.xml
    assert "w:tblBorders" in tbl_xml or "w:tcBorders" in tbl_xml


def test_render_table_from_ast_node():
    """Проверка создания таблицы из узла ASTNode."""
    doc = docx.Document()
    node = ASTNode(
        id="tbl_node_1",
        type=NodeType.TABLE,
        caption="Таблица 2.1 — Сравнительный анализ",
        headers=["Параметр", "Значение"],
        rows=[["Эффективность", "Высокая"]]
    )

    caption_p, table = render_table(doc, node=node)
    assert caption_p is not None
    assert caption_p.text == "Таблица 2.1 — Сравнительный анализ"
    assert len(table.rows) == 2
