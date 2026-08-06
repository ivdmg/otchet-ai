import pytest
import docx
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

from app.docx_engine.handlers.title_page import render_title_page
from app.docx_engine.compiler import DocxCompiler
from app.schemas.ast import DocumentASTGraph


def test_render_title_page_coursework():
    """Проверка генерации титульного листа для Курсовой работы."""
    doc = docx.Document()
    paragraphs = render_title_page(
        doc,
        doc_type="coursework",
        title="Разработка ИИ-компилятора документов",
        author_name="Иванов И.И.",
        institution="МГУ им. М.В. Ломоносова",
        city="Москва",
        year=2026
    )

    assert len(paragraphs) > 0
    full_text = "\n".join([p.text for p in doc.paragraphs])

    assert "МИНИСТЕРСТВО НАУКИ И ВЫСШЕГО ОБРАЗОВАНИЯ" in full_text
    assert "КУРСОВАЯ РАБОТА" in full_text
    assert "на тему: «Разработка ИИ-компилятора документов»" in full_text
    assert "Иванов И.И." in full_text
    assert "Москва — 2026" in full_text


def test_render_title_page_abstract():
    """Проверка генерации титульного листа для Реферата."""
    doc = docx.Document()
    render_title_page(
        doc,
        doc_type="abstract",
        title="История развития ИИ",
        author_name="Петров П.П."
    )

    full_text = "\n".join([p.text for p in doc.paragraphs])
    assert "РЕФЕРАТ" in full_text
    assert "на тему: «История развития ИИ»" in full_text


def test_render_title_page_vkr():
    """Проверка генерации титульного листа для ВКР (Дипломной работы)."""
    doc = docx.Document()
    render_title_page(
        doc,
        doc_type="vkr",
        title="Исследование нейронных сетей",
        author_name="Сидоров С.С."
    )

    full_text = "\n".join([p.text for p in doc.paragraphs])
    assert "ВЫПУСКНАЯ КВАЛИФИКАЦИОННАЯ РАБОТА" in full_text
    assert "на тему: «Исследование нейронных сетей»" in full_text


def test_compiler_title_page_integration():
    """Проверка вызова сборщика титульного листа в DocxCompiler."""
    ast = DocumentASTGraph(
        doc_type="coursework",
        title="Тестовая курсовая работа",
        author_name="Тестов Т.Т.",
        city="Санкт-Петербург",
        year=2026
    )
    compiler = DocxCompiler(ast=ast)
    doc = compiler.compile()

    full_text = "\n".join([p.text for p in doc.paragraphs])
    assert "КУРСОВАЯ РАБОТА" in full_text
    assert "Тестовая курсовая работа" in full_text
    assert "Санкт-Петербург — 2026" in full_text
